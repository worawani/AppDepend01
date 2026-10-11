"""Excel and HTML writers."""

from __future__ import annotations

import html
import json
from datetime import datetime
from pathlib import Path

from .model import KNOWN, UNKNOWN_INTERNAL, Graph

STATUS_FILLS = {
    "RED": "F8CBAD", "AMBER": "FFE699", "GREEN": "C6EFCE", "SPLIT ACROSS WAVES": "F8CBAD", "OK": "C6EFCE",
    "High": "F8CBAD", "Medium": "FFE699", "Low": "DDEBF7", "Info": "EDEDED", "no": "F8CBAD", "NO": "F8CBAD",
}
COLORED_COLS = {"Status", "Risk", "Collected", "Max Sensitivity", "Latency Sensitivity"}


# ---------------------------------------------------------------- Excel


def _write_sheet(wb, title: str, rows: list[dict], note: str = "") -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title[:31])
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="595959")
        start = 3
    if not rows:
        ws.cell(row=start, column=1, value="(none)")
        return
    cols = list(rows[0].keys())
    head_fill = PatternFill("solid", fgColor="1F3864")
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [len(c) for c in cols]
    for i, r in enumerate(rows, start + 1):
        for j, c in enumerate(cols, 1):
            v = r.get(c, "")
            cell = ws.cell(row=i, column=j, value=v)
            if c in COLORED_COLS and str(v) in STATUS_FILLS:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILLS[str(v)])
            widths[j - 1] = max(widths[j - 1], min(60, len(str(v))))
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(8, w + 2)
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(cols))}{start + len(rows)}"


def write_excel(path: Path, sheets: list[tuple[str, list[dict], str]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows, note in sheets:
        _write_sheet(wb, title, rows, note)
    wb.save(path)


# ---------------------------------------------------------------- graph JSON

WAVE_COLORS = ["#2E75B6", "#70AD47", "#ED7D31", "#7030A0", "#00B0A0", "#C00000", "#BF8F00", "#5B9BD5",
               "#A9D18E", "#F4B183", "#8064A2", "#4BACC6"]


def graph_json(g: Graph, node_filter: set[str] | None = None) -> dict:
    waves = g.waves()
    color = {w: WAVE_COLORS[i % len(WAVE_COLORS)] for i, w in enumerate(waves)}
    used = set()
    edges = []
    for e in g.edges.values():
        if e.failed_only:
            continue
        if node_filter is not None and e.client not in node_filter and e.server not in node_filter:
            continue
        used.update((e.client, e.server))
        edges.append({"from": e.client, "to": e.server, "port": e.port, "service": e.service, "cat": e.category,
                      "sens": e.sensitivity, "pers": round(e.persistence * 100, 1)})
    nodes = []
    for nid in used:
        n = g.nodes[nid]
        if n.is_vm:
            c = color.get(n.wave, "#A6A6A6")
            shape = "box" if n.collected else "ellipse"
        elif n.kind == UNKNOWN_INTERNAL:
            c, shape = "#FF0000", "diamond"
        elif n.kind == KNOWN:
            c, shape = "#7F6000", "hexagon"
        else:
            c, shape = "#404040", "triangle"
        nodes.append({"id": nid, "label": n.label, "wave": n.wave or "", "kind": n.kind, "color": c, "shape": shape,
                      "shared": n.is_shared(g.settings), "collected": n.collected,
                      "selected": bool(node_filter and nid in node_filter),
                      "app": n.vm.application if n.vm else "", "ips": sorted(n.ips)[:6]})
    return {"waves": waves, "waveColors": color, "nodes": nodes, "edges": edges}


# ---------------------------------------------------------------- HTML

_CSS = """
:root{--bg:#f6f7f9;--fg:#1d1d1f;--muted:#666;--card:#fff;--line:#e3e5e8;--head:#1F3864}
@media (prefers-color-scheme:dark){:root{--bg:#16181c;--fg:#e8e8e8;--muted:#9aa0a6;--card:#202328;--line:#33373d;--head:#2c4a7c}}
*{box-sizing:border-box}body{margin:0;font:14px/1.45 -apple-system,Segoe UI,Roboto,sans-serif;background:var(--bg);color:var(--fg)}
header{background:var(--head);color:#fff;padding:18px 24px}header h1{margin:0;font-size:20px}header p{margin:4px 0 0;opacity:.8}
main{padding:16px 24px;max-width:1600px;margin:auto}section{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:16px;margin:0 0 16px}
h2{font-size:16px;margin:0 0 10px}.cards{display:grid;grid-template-columns:repeat(auto-fill,minmax(140px,1fr));gap:12px}.card{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:12px}
.card b{display:block;font-size:24px}.card span{color:var(--muted);font-size:12px}
.tw{overflow-x:auto;max-height:520px;overflow-y:auto}table{border-collapse:collapse;width:100%;font-size:12.5px}
th,td{border-bottom:1px solid var(--line);padding:5px 8px;text-align:left;vertical-align:top;white-space:nowrap}th{position:sticky;top:0;background:var(--card);z-index:1}
td.wrap{white-space:normal;min-width:240px}.s-RED,.s-High,.s-SPLIT,.s-no,.s-NO{background:#f8cbad;color:#000}.s-AMBER,.s-Medium{background:#ffe699;color:#000}
.s-GREEN,.s-OK{background:#c6efce;color:#000}.s-Low{background:#ddebf7;color:#000}
#net{height:640px;border:1px solid var(--line);border-radius:6px;background:var(--card)}.ctl{display:flex;flex-wrap:wrap;gap:12px;align-items:center;margin-bottom:8px}
.warn{color:#c55a11}.legend span{display:inline-block;margin-right:12px;font-size:12px}.legend i{display:inline-block;width:10px;height:10px;margin-right:4px;border-radius:2px}
input[type=search]{padding:4px 8px;border:1px solid var(--line);border-radius:4px;background:var(--card);color:var(--fg)}
"""

_JS = """
const G = JSON.parse(document.getElementById('graph-data').textContent);
function draw(){
  if (typeof vis === 'undefined') { document.getElementById('net').innerHTML = '<p style="padding:16px">The graph library could not be loaded (no internet access?). The tables and the Excel workbook have the same data.</p>'; return; }
  const wave = document.getElementById('f-wave').value, showInfra = document.getElementById('f-infra').checked,
        showAdmin = document.getElementById('f-admin').checked, q = document.getElementById('f-q').value.toLowerCase();
  let edges = G.edges.filter(e => e.cat === 'app' || (showInfra && e.cat === 'infra') || (showAdmin && e.cat === 'admin'));
  const byId = Object.fromEntries(G.nodes.map(n => [n.id, n]));
  if (wave) edges = edges.filter(e => byId[e.from].wave === wave || byId[e.to].wave === wave);
  if (q) edges = edges.filter(e => e.from.toLowerCase().includes(q) || e.to.toLowerCase().includes(q));
  const used = new Set(); edges.forEach(e => { used.add(e.from); used.add(e.to); });
  const nodes = G.nodes.filter(n => used.has(n.id)).map(n => ({id:n.id, label:n.label + (n.wave ? '\\n' + n.wave : ''), shape:n.shape,
     color:{background:n.color, border:n.selected ? '#000' : n.color}, borderWidth:n.selected ? 3 : 1,
     font:{color:(n.shape==='box') ? '#fff' : (matchMedia('(prefers-color-scheme: dark)').matches ? '#eee' : '#222'), size:12},
     title:`${n.id}\\nWave: ${n.wave || '-'}\\nApp: ${n.app || '-'}\\nIPs: ${n.ips.join(', ')}\\n${n.collected ? 'collected' : (n.kind==='vm' ? 'NOT collected' : n.kind)}${n.shared ? '\\nshared service' : ''}`}));
  const vedges = edges.map((e, i) => ({id:i, from:e.from, to:e.to, arrows:'to', label:String(e.port),
     title:`${e.from} -> ${e.to}:${e.port} (${e.service})\\n${e.cat}, latency ${e.sens}, persistence ${e.pers}%`,
     dashes:e.cat !== 'app', width: e.sens === 'High' ? 2.5 : 1, font:{size:9, align:'middle'},
     color:{color: byId[e.from].wave && byId[e.to].wave && byId[e.from].wave !== byId[e.to].wave ? '#C00000' : (e.cat==='app' ? '#7f7f7f' : '#bfbfbf')}}));
  document.getElementById('net-count').textContent = `${nodes.length} nodes, ${vedges.length} edges`;
  new vis.Network(document.getElementById('net'), {nodes:new vis.DataSet(nodes), edges:new vis.DataSet(vedges)},
     {physics:{solver:'forceAtlas2Based', stabilization:{iterations:250}}, interaction:{hover:true, navigationButtons:true}, layout:{improvedLayout: nodes.length < 150}});
}
function initGraph(){
  const sel = document.getElementById('f-wave');
  G.waves.forEach(w => { const o = document.createElement('option'); o.value = o.textContent = w; sel.appendChild(o); });
  document.getElementById('legend').innerHTML = G.waves.map(w => `<span><i style="background:${G.waveColors[w]}"></i>${w}</span>`).join('') +
    '<span><i style="background:#A6A6A6"></i>no wave</span><span><i style="background:#FF0000"></i>unknown IP</span><span><i style="background:#7F6000"></i>known endpoint</span><span><i style="background:#404040"></i>internet</span><span>red edge = cross-wave</span><span>box = collected, ellipse = no data</span>';
  ['f-wave','f-infra','f-admin'].forEach(id => document.getElementById(id).addEventListener('change', draw));
  document.getElementById('f-q').addEventListener('input', () => { clearTimeout(window._t); window._t = setTimeout(draw, 300); });
  draw();
}
document.querySelectorAll('input[data-filter]').forEach(inp => inp.addEventListener('input', () => {
  const q = inp.value.toLowerCase();
  document.querySelectorAll('#' + inp.dataset.filter + ' tbody tr').forEach(tr => { tr.style.display = tr.textContent.toLowerCase().includes(q) ? '' : 'none'; });
}));
initGraph();
"""

VIS_CDN = "https://unpkg.com/vis-network@9.1.9/standalone/umd/vis-network.min.js"


def _vis_script() -> str:
    """Inline the vendored vis-network so the report works on offline jump hosts; CDN as fallback."""
    vendored = Path(__file__).with_name("vendor") / "vis-network.min.js"
    if vendored.exists():
        return "<script>" + vendored.read_text(encoding="utf-8").replace("</script", "<\\/script") + "</script>"
    return f'<script src="{VIS_CDN}"></script>'


WRAP_COLS = {"Issues", "Members", "Impact", "Verdict", "Action", "Peers", "Processes", "Suggestion", "Note", "VMs Without Data", "Linked To"}


def html_table(tid: str, rows: list[dict], limit: int = 1000, cols: list[str] | None = None) -> str:
    if not rows:
        return "<p>(none)</p>"
    cols = cols or list(rows[0].keys())
    out = [f'<input type="search" placeholder="Filter…" data-filter="{tid}"> ']
    if len(rows) > limit:
        out.append(f"<span class='warn'>Showing the first {limit} of {len(rows)} rows. The Excel workbook has all of them.</span>")
    out.append(f'<div class="tw"><table id="{tid}"><thead><tr>' + "".join(f"<th>{html.escape(c)}</th>" for c in cols) + "</tr></thead><tbody>")
    for r in rows[:limit]:
        cells = []
        for c in cols:
            v = str(r.get(c, ""))
            cls = []
            if c in COLORED_COLS and v:
                cls.append("s-" + v.split()[0])
            if c in WRAP_COLS:
                cls.append("wrap")
            cells.append(f'<td class="{" ".join(cls)}">{html.escape(v)}</td>' if cls else f"<td>{html.escape(v)}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</tbody></table></div>")
    return "".join(out)


def write_html(path: Path, title: str, subtitle: str, cards: list[tuple[str, str]], sections: list[tuple[str, str]],
               graph: dict, warnings: list[str]) -> None:
    card_html = "".join(f'<div class="card"><b>{html.escape(str(v))}</b><span>{html.escape(k)}</span></div>' for k, v in cards)
    warn_html = ""
    if warnings:
        items = "".join(f"<li>{html.escape(w)}</li>" for w in warnings[:200])
        more = f"<li>… and {len(warnings) - 200} more</li>" if len(warnings) > 200 else ""
        warn_html = f'<section><h2 class="warn">Data quality warnings ({len(warnings)})</h2><ul>{items}{more}</ul></section>'
    body = "".join(f"<section><h2>{html.escape(h)}</h2>{content}</section>" for h, content in sections)
    graph_section = """<section><h2>Dependency graph</h2>
<div class="ctl"><label>Wave <select id="f-wave"><option value="">All</option></select></label>
<label><input type="checkbox" id="f-infra"> show infra flows</label><label><input type="checkbox" id="f-admin"> show admin flows</label>
<input type="search" id="f-q" placeholder="Find node…"><span id="net-count" style="color:var(--muted)"></span></div>
<div id="legend" class="legend"></div><div id="net"></div></section>"""
    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title><style>{_CSS}</style>
{_vis_script()}</head><body>
<header><h1>{html.escape(title)}</h1><p>{html.escape(subtitle)}</p></header><main>
<div class="cards" style="margin-bottom:16px">{card_html}</div>{warn_html}{graph_section}{body}</main>
<script type="application/json" id="graph-data">{json.dumps(graph).replace("</", "<\\/")}</script>
<script>{_JS}</script></body></html>"""
    path.write_text(doc, encoding="utf-8")


def stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")

