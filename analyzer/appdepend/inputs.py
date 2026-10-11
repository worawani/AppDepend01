"""Loaders for the analyzer's inputs: VM inventory (RVTools or plain CSV),
wave plan, known endpoints, and collector output."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from .netutil import CidrTable, split_ips

# ---------------------------------------------------------------- tabular reading


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (h or "").lower())


def read_table(path: Path, sheet: str | None = None) -> list[dict[str, str]]:
    """Rows of a CSV or one sheet of an .xlsx as dicts with the original headers."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip() if h is not None else "" for h in next(rows)]
        except StopIteration:
            return []
        out = []
        for r in rows:
            if r is None or all(v is None for v in r):
                continue
            out.append({header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header)})
        wb.close()
        return out
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(f, dialect=dialect)]


def pick(row: dict[str, str], *aliases: str) -> str:
    """First non-empty value whose header matches any alias (case/punctuation-insensitive)."""
    normed = {_norm(k): v for k, v in row.items()}
    for a in aliases:
        v = normed.get(_norm(a))
        if v:
            return v
    return ""


# ---------------------------------------------------------------- inventory


@dataclass
class VM:
    name: str
    ips: set[str] = field(default_factory=set)
    dns_name: str = ""
    os: str = ""
    cluster: str = ""
    application: str = ""
    owner: str = ""
    role: str = ""              # 'infra'/'shared' marks a shared-service VM
    wave: str = ""
    power: str = ""
    cpus: str = ""
    memory_mb: str = ""
    in_inventory: bool = True

    @property
    def is_shared(self) -> bool:
        return self.role.lower() in ("infra", "infrastructure", "shared", "shared-service", "shared service")

    @property
    def is_windows(self) -> bool:
        return "windows" in self.os.lower()


NAME_COLS = ("VM", "VM Name", "VMName", "Name", "Server", "Server Name", "Hostname")
IP_COLS = ("IPs", "IP Addresses", "IPAddresses", "IP Address", "IP", "Primary IP Address", "IPv4 Address")


def _vm_from_row(row: dict[str, str]) -> VM | None:
    name = pick(row, *NAME_COLS)
    if not name or pick(row, "Template").lower() == "true":
        return None
    vm = VM(name=name)
    for col in IP_COLS:
        vm.ips.update(split_ips(pick(row, col)))
    vm.dns_name = pick(row, "DNS Name", "FQDN", "DNSName")
    vm.os = pick(row, "OS", "OS according to the VMware Tools", "OS according to the configuration file", "Guest OS")
    vm.cluster = pick(row, "Cluster")
    vm.application = pick(row, "Application", "App", "Application Name", "Service")
    vm.owner = pick(row, "Owner", "App Owner", "Application Owner")
    vm.role = pick(row, "Role", "Tier", "Type")
    vm.wave = pick(row, "Wave", "Migration Wave")
    vm.power = pick(row, "Powerstate", "Power State", "PowerState")
    vm.cpus = pick(row, "CPUs", "vCPU", "Num CPUs", "CPU")
    vm.memory_mb = pick(row, "Memory", "Memory MB", "MemoryMB", "RAM")
    return vm


def load_inventory(paths: list[Path]) -> dict[str, VM]:
    """Merge one or more inventory files. RVTools .xlsx (vInfo + vNetwork) is detected automatically."""
    vms: dict[str, VM] = {}

    def merge(vm: VM) -> None:
        key = vm.name.lower()
        cur = vms.get(key)
        if cur is None:
            vms[key] = vm
            return
        cur.ips |= vm.ips
        for attr in ("dns_name", "os", "cluster", "application", "owner", "role", "wave", "power", "cpus", "memory_mb"):
            if not getattr(cur, attr) and getattr(vm, attr):
                setattr(cur, attr, getattr(vm, attr))

    for path in paths:
        sheets: list[str | None] = [None]
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            names = load_workbook(path, read_only=True).sheetnames
            if "vInfo" in names:  # RVTools export
                sheets = ["vInfo"] + [s for s in ("vNetwork",) if s in names]
        for sheet in sheets:
            for row in read_table(path, sheet):
                vm = _vm_from_row(row)
                if vm:
                    merge(vm)
    return vms


def load_waves(path: Path) -> dict[str, dict[str, str]]:
    """VM -> {'wave': ..., 'date': ..., 'application': ...}."""
    out = {}
    for row in read_table(path):
        name = pick(row, *NAME_COLS)
        wave = pick(row, "Wave", "Migration Wave", "Move Group")
        if name and wave:
            out[name.lower()] = {
                "wave": wave,
                "date": pick(row, "Date", "Cutover Date", "Cutover", "Migration Date"),
                "application": pick(row, "Application", "App"),
            }
    return out


@dataclass
class KnownEndpoint:
    cidr: str
    label: str
    category: str   # app | infra | admin | external
    note: str = ""


def load_known(path: Path | None) -> CidrTable:
    table = CidrTable()
    if not path:
        return table
    for row in read_table(path):
        cidr = pick(row, "CIDR", "IP", "Network", "Subnet")
        if not cidr:
            continue
        cat = (pick(row, "Category", "Type", "Role") or "app").lower()
        if cat in ("shared", "infrastructure"):
            cat = "infra"
        table.add(cidr, KnownEndpoint(cidr, pick(row, "Label", "Name", "Description") or cidr, cat, pick(row, "Note", "Notes")))
    return table


def natural_key(text: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", text or "")]


# ---------------------------------------------------------------- collector output


@dataclass
class Flow:
    direction: str          # Inbound | Outbound
    remote: str
    port: int
    process: str
    process_path: str
    services: str
    states: set[str]
    first_seen: str
    last_seen: str
    samples: int
    max_concurrent: int


@dataclass
class HostData:
    hostname: str
    fqdn: str = ""
    os: str = ""
    ips: set[str] = field(default_factory=set)
    start: str = ""
    end: str = ""
    samples: int = 0
    interval: int = 0
    flows: dict[tuple, Flow] = field(default_factory=dict)
    listeners: list[dict[str, str]] = field(default_factory=list)
    services: list[dict[str, str]] = field(default_factory=list)
    source_dirs: set[str] = field(default_factory=set)

    @property
    def days(self) -> float:
        try:
            s = datetime.fromisoformat(self.start)
            e = datetime.fromisoformat(self.end)
        except (TypeError, ValueError):
            return 0.0
        return max(0.0, (e - s).total_seconds() / 86400)


def _int(v, default=0) -> int:
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def load_collected(root: Path) -> dict[str, HostData]:
    """Recursively load every <HOST>_*.csv set under root. Duplicate copies of the same
    host (for example harvested twice) are merged with max() so nothing is double-counted."""
    hosts: dict[str, HostData] = {}
    prefixes: dict[tuple[Path, str], None] = {}
    for f in root.rglob("*.csv"):
        m = re.match(r"(.+)_(connections|listeners|services|hostinfo)\.csv$", f.name, re.I)
        if m:
            prefixes[(f.parent, m.group(1))] = None

    for (folder, prefix) in prefixes:
        def path(kind: str) -> Path:
            return folder / f"{prefix}_{kind}.csv"

        info_rows = read_table(path("hostinfo")) if path("hostinfo").exists() else []
        info = info_rows[0] if info_rows else {}
        hostname = (info.get("Hostname") or prefix).upper()
        h = hosts.setdefault(hostname.lower(), HostData(hostname=hostname))
        h.source_dirs.add(str(folder))
        if info:
            h.fqdn = h.fqdn or info.get("FQDN", "")
            h.os = h.os or info.get("OS", "")
            h.ips.update(split_ips(info.get("IPAddresses", "")))
            if info.get("CollectionStart") and (not h.start or info["CollectionStart"] < h.start):
                h.start = info["CollectionStart"]
            if info.get("CollectionEnd") and info["CollectionEnd"] > h.end:
                h.end = info["CollectionEnd"]
            h.samples = max(h.samples, _int(info.get("Samples")))
            h.interval = h.interval or _int(info.get("IntervalSeconds"))

        if path("connections").exists():
            for r in read_table(path("connections")):
                port = _int(r.get("Port"), -1)
                remote = r.get("RemoteAddress", "")
                direction = r.get("Direction", "")
                if port < 0 or not remote or direction not in ("Inbound", "Outbound"):
                    continue
                key = (direction, remote, port, r.get("ProcessName", ""))
                fl = Flow(direction, remote, port, r.get("ProcessName", ""), r.get("ProcessPath", ""),
                          r.get("ServiceNames", ""), set(filter(None, r.get("States", "").split(";"))),
                          r.get("FirstSeen", ""), r.get("LastSeen", ""), _int(r.get("SampleCount")),
                          _int(r.get("MaxConcurrent")))
                cur = h.flows.get(key)
                if cur is None:
                    h.flows[key] = fl
                else:
                    cur.states |= fl.states
                    cur.samples = max(cur.samples, fl.samples)
                    cur.max_concurrent = max(cur.max_concurrent, fl.max_concurrent)
                    cur.first_seen = min(filter(None, [cur.first_seen, fl.first_seen]), default="")
                    cur.last_seen = max(cur.last_seen, fl.last_seen)

        if path("listeners").exists():
            seen = {(l.get("LocalAddress"), l.get("LocalPort")) for l in h.listeners}
            h.listeners += [r for r in read_table(path("listeners")) if (r.get("LocalAddress"), r.get("LocalPort")) not in seen]
        if path("services").exists():
            seen_s = {s.get("Name") for s in h.services}
            h.services += [r for r in read_table(path("services")) if r.get("Name") not in seen_s]
    return hosts
